"""Build the committed knowledge files from the local organizer materials.

Run with `uv run afs-assist kb build`. Inputs live in resources/ (gitignored);
outputs in knowledge/ (committed). The build is deterministic: same inputs,
same bytes out.
"""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import pymupdf
import yaml
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from afs_assistant import paths
from afs_assistant.normalize import normalize

DOCX_NAME = "AFS_Examples_and_Reference.docx"
S06_NAME = "S06_Laboratory_Link.txt"
# Sections of the organizer reference docx that may be cited as REF §n (user decision 2026-10-17).
REF_SECTIONS = {
    3: "3. Answer structure and good practice",
    5: "5. Method overview",
}
REF_SECTION_END = {3: "4. AFS sources and citations", 5: None}


def sha256(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def extract_pdf_pages(doc_id: str) -> list[dict]:
    pdf = pymupdf.open(paths.MATERIALS / paths.SOURCE_PDFS[doc_id])
    return [
        {"doc": doc_id, "page": i + 1, "text": normalize(page.get_text()),
         "extraction": f"pymupdf-{pymupdf.VersionBind}"}
        for i, page in enumerate(pdf)
    ]


def _docx_blocks(document: Document) -> list[str]:
    """Paragraphs and tables in body order; table rows rendered as 'a | b | c'."""
    blocks = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            text = Paragraph(child, document).text.strip()
            if text:
                blocks.append(text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, document).rows:
                blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    return blocks


def extract_ref_sections() -> list[dict]:
    blocks = _docx_blocks(Document(paths.MATERIALS / DOCX_NAME))
    pages = []
    for number, heading in REF_SECTIONS.items():
        start = blocks.index(heading)
        end_heading = REF_SECTION_END[number]
        end = blocks.index(end_heading) if end_heading else len(blocks)
        pages.append({"doc": "REF", "page": number, "text": normalize("\n".join(blocks[start:end])),
                      "extraction": "python-docx"})
    return pages


def extract_s06() -> list[dict]:
    text = (paths.MATERIALS / S06_NAME).read_text(encoding="utf-8")
    return [{"doc": "S06", "page": 1, "text": normalize(text), "extraction": "plain-text"}]


_SLIDE_HEADING = re.compile(r"^## Slide (\d+)\b.*$", re.MULTILINE)


def parse_s02_notes(markdown: str) -> list[dict]:
    """Split our S02 transcription notes into one page record per slide."""
    matches = list(_SLIDE_HEADING.finditer(markdown))
    pages = []
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(markdown)
        body = markdown[m.end():end]
        pages.append({"doc": "S02", "page": int(m.group(1)), "text": normalize(body),
                      "extraction": "manual-transcription"})
    return pages


def load_scopes() -> dict[tuple[str, int], dict]:
    if not paths.PAGE_SCOPES.exists():
        return {}
    data = yaml.safe_load(paths.PAGE_SCOPES.read_text(encoding="utf-8")) or {}
    scopes = {}
    for doc, entries in data.items():
        for page, entry in entries.items():
            scopes[(doc, int(page))] = entry
    return scopes


def build_pages() -> list[dict]:
    pages = []
    for doc_id in ("S01", "S03", "S04", "S05"):
        pages += extract_pdf_pages(doc_id)
    if paths.S02_NOTES.exists():
        pages += parse_s02_notes(paths.S02_NOTES.read_text(encoding="utf-8"))
    pages += extract_s06()
    pages += extract_ref_sections()

    scopes = load_scopes()
    order = {"S01": 1, "S02": 2, "S03": 3, "S04": 4, "S05": 5, "S06": 6, "REF": 7}
    pages.sort(key=lambda p: (order[p["doc"]], p["page"]))
    records = []
    for p in pages:
        meta = scopes.get((p["doc"], p["page"]), {})
        records.append({
            "doc": p["doc"], "page": p["page"],
            "title": meta.get("title", ""),
            "scope": meta.get("scope", []),
            "text": p["text"], "sha256": sha256(p["text"]),
            "extraction": p["extraction"],
        })
    return records


XLSX_NAME = "AFS_method_comparison_table.xlsx"
_PATH_RANGE = re.compile(r"^\s*([\d.]+)\s*(?:-\s*([\d.]+))?\s*(um|mm)\s*$")
_TO_MM = {"um": 1e-3, "mm": 1.0}


def parse_path_mm(value: str) -> dict:
    """'1-6 um' -> {'min_mm': 0.001, 'max_mm': 0.006}; 'any' -> nulls."""
    m = _PATH_RANGE.match(value or "")
    if not m:
        return {"min_mm": None, "max_mm": None}
    lo = float(m.group(1)) * _TO_MM[m.group(3)]
    hi = float(m.group(2) or m.group(1)) * _TO_MM[m.group(3)]
    return {"min_mm": round(lo, 6), "max_mm": round(hi, 6)}


def build_table1() -> list[dict]:
    import openpyxl

    sheet = openpyxl.load_workbook(paths.MATERIALS / XLSX_NAME, read_only=True)["method_comparison"]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    records = []
    for i, row in enumerate(rows[1:], start=1):
        rec = {h: (str(v).strip() if v is not None else "") for h, v in zip(header, row)}
        rec["sample_types"] = [s.strip() for s in rec["sample_type"].split(",") if s.strip()]
        rec.update(parse_path_mm(rec["optical_path_or_depth"]))
        rec["row"] = i
        rec["citation"] = "S01 Table 1 (p.6)"
        records.append(rec)
    return records


def write_jsonl(records: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for r in records:
            fh.write(json.dumps(r, ensure_ascii=False, sort_keys=True) + "\n")


def build_all() -> dict[str, int]:
    pages = build_pages()
    write_jsonl(pages, paths.PAGES_FILE)
    table1 = build_table1()
    paths.TABLE1_FILE.write_text(
        json.dumps({"_about": "S01 Table 1 (p.6), from the organizer xlsx transcription. "
                    "sensitivity/selectivity/cost are the source's qualitative ratings.",
                    "rows": table1}, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
        encoding="utf-8")
    counts: dict[str, int] = {}
    for p in pages:
        counts[p["doc"]] = counts.get(p["doc"], 0) + 1
    return counts
